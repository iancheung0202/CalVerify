import os
import logging
import polars as pl
import time
import asyncio
import uuid
import uvicorn
import json
import re

from fastapi import FastAPI, Request, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, JSONResponse

from pathlib import Path
from typing import Dict, List, Any
from dotenv import load_dotenv
from datetime import datetime, timezone
from collections import defaultdict
from asyncio import Lock
from google.oauth2.service_account import Credentials
from google.oauth2 import id_token
from google.auth.transport import requests as google_requests
from googleapiclient.discovery import build

from bot import init_bot, get_user, get_profile, assign_role

load_dotenv()

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(title="Discord Verification Student Search")

CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID", "")
SHEET_ID = os.getenv("VERIFICATION_FORM_SHEET_ID")

MAX_CONTENT_LENGTH = 10 * 1024 * 1024  # 10MB

sheets, sid = None, None

def init_sheets():
    global sheets, sid
    try:
        creds_file = Path(__file__).parent / "service-account-key.json"
        if creds_file.exists():
            creds = Credentials.from_service_account_file(creds_file, scopes=['https://www.googleapis.com/auth/spreadsheets'])
            sheets = build('sheets', 'v4', credentials=creds)
            logger.info("Google Sheets API initialized successfully")
            try:
                spreadsheet = sheets.spreadsheets().get(spreadsheetId=SHEET_ID).execute()
                for s in spreadsheet.get('sheets', []):
                    if s['properties']['title'] == 'Form Responses 1':
                        sid = s['properties']['sheetId']
                        logger.info(f"Found sheet 'Form Responses 1' with ID: {sid}")
                        break
                if sid is None:
                    logger.warning("Could not find 'Form Responses 1' sheet, defaulting to first sheet")
                    if spreadsheet.get('sheets'):
                        sid = spreadsheet['sheets'][0]['properties']['sheetId']
            except Exception as e:
                logger.warning(f"Could not fetch sheet ID: {str(e)}, defaulting to 0")
                sid = 0
        else:
            logger.warning(f"Service account credentials not found at {creds_file}")
    except Exception as e:
        logger.error(f"Error initializing Google Sheets API: {str(e)}")


app.add_middleware(
    CORSMiddleware,
    allow_origins=["https://verify.iancheung.dev", "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["POST", "GET"],
    allow_headers=["Content-Type"],
)

@app.middleware("http")
async def limit_request_body_size(request: Request, call_next):
    content_length = request.headers.get("content-length")
    if content_length and int(content_length) > MAX_CONTENT_LENGTH:
        return JSONResponse(status_code=413, content={"detail": "Request too large"})
    return await call_next(request)

@app.middleware("http")
async def add_security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["X-XSS-Protection"] = "0"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    if request.url.scheme == "https":
        response.headers["Strict-Transport-Security"] = "max-age=63072000; includeSubDomains; preload"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; "
        "script-src 'self' https://accounts.google.com https://cdn.jsdelivr.net https://cdnjs.cloudflare.com 'unsafe-inline'; "
        "style-src 'self' https://cdn.jsdelivr.net https://cdnjs.cloudflare.com 'unsafe-inline'; "
        "img-src 'self' https: data:; "
        "font-src 'self' https://cdnjs.cloudflare.com data:; "
        "connect-src 'self'; "
        "frame-src https://accounts.google.com; "
        "object-src 'none'; "
        "base-uri 'self'; "
        "form-action 'self'"
    )
    return response

def verify_ref(request: Request) -> bool:
    referer = request.headers.get("referer", "")
    if not referer:
        return False
    return "verify.iancheung.dev" in referer

login_fails: Dict[str, List[float]] = defaultdict(list)
LOGIN_MAX = 5
LOGIN_WIN = 300
LOGIN_LOCKOUT = 900

searches: Dict[str, List[float]] = defaultdict(list)
SEARCH_MAX = 120
SEARCH_WIN = 60

SESSION_TTL = 3600
SESSION_COOKIE = "session_token"
SESSION_MAX_AGE = 3600

sessions: Dict[str, Dict[str, Any]] = {}
sessions_lock = Lock()

def client_ip(request: Request) -> str:
    forwarded = request.headers.get("x-forwarded-for")
    if forwarded:
        return forwarded.split(",")[0].strip()
    real = request.headers.get("x-real-ip")
    if real:
        return real.strip()
    return request.client.host if request.client else "unknown"

def rate_limited(ip: str, reqs: Dict[str, List[float]], limit: int, window: int) -> bool:
    now = time.time()
    reqs[ip] = [t for t in reqs[ip] if now - t < window]
    if len(reqs[ip]) >= limit:
        return True
    reqs[ip].append(now)
    return False

def login_locked(ip: str) -> bool:
    now = time.time()
    login_fails[ip] = [t for t in login_fails[ip] if now - t < LOGIN_WIN]
    if len(login_fails[ip]) >= LOGIN_MAX:
        if now - min(login_fails[ip]) < LOGIN_LOCKOUT:
            return True
    return False

async def authed(request: Request) -> bool:
    token = request.cookies.get(SESSION_COOKIE)
    if not token:
        return False
    data = sessions.get(token)
    if not data:
        return False
    if time.time() > data.get("expires_at", 0):
        sessions.pop(token, None)
        return False
    return True

def create_session(ip: str, email: str) -> str:
    token = str(uuid.uuid4())
    expires = time.time() + SESSION_TTL
    sessions[token] = {"expires_at": expires, "ip": ip, "email": email, "created_at": time.time()}
    logger.info(f"Created session token for IP: {ip}, Email: {email}")
    return token

async def cleanup_sessions():
    now = time.time()
    async with sessions_lock:
        expired = [t for t, d in list(sessions.items()) if now > d.get("expires_at", 0)]
        for t in expired:
            sessions.pop(t, None)
        if expired:
            logger.info(f"Cleaned up {len(expired)} expired sessions")

dfs: Dict[str, pl.DataFrame] = {}
meta: Dict[str, Dict[str, Any]] = {}

def load_csvs():
    calverify_dir = Path(__file__).parent / "data"
    csv_files = list(calverify_dir.glob("*.csv"))
    logger.info(f"Found {len(csv_files)} CSV files")
    for csv_file in csv_files:
        try:
            df = pl.read_csv(csv_file)
            name = csv_file.name
            dfs[name] = df
            mod = os.path.getmtime(csv_file)
            mod_dt = datetime.fromtimestamp(mod, tz=timezone.utc)
            meta[name] = {"last_modified_iso": mod_dt.isoformat(), "total_rows": len(df)}
            logger.info(f"Loaded {name} with {len(df)} rows (modified: {mod_dt.isoformat()})")
        except Exception as e:
            logger.error(f"Error loading {csv_file.name}: {str(e)}")
    logger.info(f"Successfully loaded {len(dfs)} DataFrames")

def search_dfs(query: str) -> Dict[str, List[Dict[str, Any]]]:
    results = {}
    qlower = query.lower()
    for name, df in dfs.items():
        cond = None
        for col in df.columns:
            c = pl.col(col).cast(pl.Utf8).str.to_lowercase().str.contains(qlower)
            cond = c if cond is None else (cond | c)
        if cond is not None:
            rows = df.filter(cond).to_dicts()
            if rows:
                results[name] = rows
    return results

def reload_csvs():
    global dfs, meta
    dfs.clear()
    meta.clear()
    load_csvs()

def load_manifest() -> Dict[str, Any]:
    path = Path(__file__).parent / "manifest.json"
    default = {"sheetFiles": {}, "authorizedEmails": []}
    try:
        if path.exists():
            with open(path, 'r') as f:
                data = json.load(f)
                if "authorizedEmails" not in data:
                    data["authorizedEmails"] = []
                return data
    except Exception as e:
        logger.error(f"Error loading manifest.json: {str(e)}")
    return default

def save_manifest(manifest: Dict[str, Any]):
    path = Path(__file__).parent / "manifest.json"
    try:
        with open(path, 'w') as f:
            json.dump(manifest, f, indent=2)
        logger.info("Manifest saved successfully")
    except Exception as e:
        logger.error(f"Error saving manifest.json: {str(e)}")

def parse_sheet_url(url: str) -> tuple:
    try:
        match = re.search(r'/spreadsheets/d/([a-zA-Z0-9-_]+)', url)
        gid_match = re.search(r'[?#&]gid=(\d+)', url)
        if match and gid_match:
            return match.group(1), gid_match.group(1)
    except Exception as e:
        logger.error(f"Error extracting sheet info from URL: {str(e)}")
    return None, None

def download_sheet(drive, ss, sheet_id: str, gid: str = "0") -> tuple:
    try:
        name = "Sheet"
        try:
            if ss:
                sp = ss.spreadsheets().get(spreadsheetId=sheet_id, fields='properties/title').execute()
                name = sp.get('properties', {}).get('title', 'Sheet')
                logger.info(f"Spreadsheet title retrieved from Sheets API: '{name}'")
            else:
                logger.warning("Sheets service not available for fetching spreadsheet title")
        except Exception as e:
            logger.warning(f"Could not fetch spreadsheet title via Sheets API: {str(e)}, using default name")
            name = "Sheet"
        try:
            fm = drive.files().get(fileId=sheet_id, fields='mimeType').execute()
            mime = fm.get('mimeType', '')
            is_gsheet = mime == 'application/vnd.google-apps.spreadsheet'
            logger.info(f"File MIME type: {mime}")
        except Exception as e:
            logger.warning(f"Could not fetch file MIME type: {str(e)}, assuming XLSX format")
            is_gsheet = False
        tmp_path = Path(__file__).parent / f"temp_sheet_{sheet_id}.xlsx"
        try:
            if is_gsheet:
                logger.info(f"Exporting native Google Sheet to XLSX format")
                request = drive.files().export(fileId=sheet_id, mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            else:
                logger.info(f"Downloading XLSX file directly")
                request = drive.files().get_media(fileId=sheet_id)
            data = request.execute()
        except Exception as e:
            logger.error(f"Failed to download/export file: {str(e)}")
            return None, None, None
        if len(data) > MAX_CONTENT_LENGTH:
            logger.error(f"Downloaded file too large: {len(data)} bytes")
            return None, None, None
        try:
            with open(tmp_path, 'wb') as f:
                f.write(data)
        except Exception as e:
            logger.error(f"Failed to write temporary file: {str(e)}")
            return None, None, None
        try:
            import openpyxl
            wb = openpyxl.load_workbook(tmp_path)
            sname = None
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                if str(getattr(ws, 'sheet_id', None)) == gid:
                    sname = sheet
                    break
            if not sname:
                logger.warning(f"Sheet with gid {gid} not found, using first sheet")
                if wb.sheetnames:
                    sname = wb.sheetnames[0]
                else:
                    logger.error(f"No sheets found in workbook")
                    if tmp_path.exists():
                        tmp_path.unlink()
                    return None, None, None
            logger.info(f"Using sheet: '{sname}' (gid: {gid})")
        except ImportError:
            logger.error("openpyxl not installed")
            if tmp_path.exists():
                tmp_path.unlink()
            return None, None, None
        except Exception as e:
            logger.error(f"Error parsing workbook: {str(e)}")
            if tmp_path.exists():
                tmp_path.unlink()
            return None, None, None
        try:
            df = pl.read_excel(tmp_path, sheet_name=sname)
        except Exception as e:
            logger.error(f"Error reading sheet '{sname}': {str(e)}")
            if tmp_path.exists():
                tmp_path.unlink()
            return None, None, None
        csv_content = df.write_csv()
        if tmp_path.exists():
            tmp_path.unlink()
        logger.info(f"Successfully downloaded sheet '{sname}' with {len(df)} rows")
        return csv_content, name, sname
    except Exception as e:
        logger.error(f"Error downloading sheet as CSV: {str(e)}")
        return None, None, None


def fetch_entries() -> List[Dict[str, Any]]:
    if not sheets:
        logger.error("Google Sheets API not initialized")
        return []
    try:
        sp_data = sheets.spreadsheets().get(
            spreadsheetId=SHEET_ID,
            includeGridData=True,
            ranges=["Form Responses 1!B:I"],
            fields="sheets/data/rowData/values(userEnteredValue,userEnteredFormat/backgroundColor)"
        ).execute()
        sd = sp_data.get('sheets', [{}])[0]
        rd = sd.get('data', [{}])[0].get('rowData', [])
        if len(rd) < 2:
            logger.warning("No student entries found in sheet")
            return []

        def rgb_hex(c):
            if not c:
                return None
            r = round((c.get('red', 0) or 0) * 255)
            g = round((c.get('green', 0) or 0) * 255)
            b = round((c.get('blue', 0) or 0) * 255)
            return f"#{r:02x}{g:02x}{b:02x}"

        sc = 0
        nc = 1
        ec = 0
        gc = 3
        dc = 4
        sic = 5
        noc = 6
        dic = 7

        entries = []
        for idx in range(len(rd) - 1, 0, -1):
            rfd = rd[idx]
            vs = rfd.get('values', [])
            row = []
            for cell in vs:
                val = cell.get('userEnteredValue', {})
                if 'stringValue' in val:
                    row.append(val['stringValue'])
                elif 'numberValue' in val:
                    row.append(str(val['numberValue']))
                elif 'boolValue' in val:
                    row.append(str(val['boolValue']))
                else:
                    row.append('')
            name = row[nc] if len(row) > nc else ""
            email = row[ec] if len(row) > ec else ""
            grad = row[gc] if len(row) > gc else ""
            disc = row[dc] if len(row) > dc else ""
            sid_val = row[sic] if len(row) > sic else ""
            notes = row[noc] if len(row) > noc else ""
            did = row[dic] if len(row) > dic else ""
            if not sid_val.strip():
                continue
            bg = None
            if sc < len(vs):
                fmt = vs[sc].get('userEnteredFormat', {})
                bgd = fmt.get('backgroundColor', {})
                bg = rgb_hex(bgd) if bgd else None
            entries.append({
                "row_index": idx,
                "name": name.strip(),
                "email": email.strip(),
                "graduating_class": grad.strip(),
                "discord_username": disc.strip(),
                "student_id": sid_val.strip(),
                "notes": notes.strip(),
                "discord_user_id": did.strip(),
                "background_color": bg,
            })
        logger.info(f"Fetched {len(entries)} student entries from Google Sheet")
        return entries
    except Exception as e:
        logger.error(f"Error fetching student entries: {str(e)}")
        return []




async def monitor_entries():
    logger.info("Initializing background entry monitor...")
    await asyncio.sleep(5)
    initial = fetch_entries()
    last_max = max([e["row_index"] for e in initial]) if initial else 0
    logger.info(f"Monitor started. Current max row: {last_max}")
    while True:
        await asyncio.sleep(60)
        try:
            entries = fetch_entries()
            if not entries:
                continue
            cur_max = max([e["row_index"] for e in entries])
            if cur_max > last_max:
                new = [e for e in entries if e["row_index"] > last_max]
                for entry in new:
                    sid_val = entry.get("student_id", "").strip()
                    if not sid_val:
                        continue
                    results = search_dfs(sid_val)
                    total = sum(len(v) for v in results.values())
                    if total >= 1:
                        logger.info(f"New submission from {entry.get('name')} ({entry.get('email')})")
                last_max = cur_max
        except Exception as e:
            logger.error(f"Error in monitor_entries loop: {e}")


@app.on_event("startup")
async def startup_event():
    load_csvs()
    init_sheets()
    asyncio.create_task(monitor_entries())
    asyncio.create_task(session_cleanup_loop())
    await init_bot()

async def session_cleanup_loop():
    while True:
        await asyncio.sleep(300)
        await cleanup_sessions()

@app.get("/config")
async def get_config():
    sheet_url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit" if SHEET_ID else None
    return {"google_client_id": CLIENT_ID, "sheet_url": sheet_url}

@app.post("/login")
async def login(request: Request):
    if not verify_ref(request):
        raise HTTPException(status_code=403, detail="Access denied")
    ip = client_ip(request)
    if login_locked(ip):
        logger.warning(f"Login attempt from locked-out IP: {ip}")
        raise HTTPException(status_code=429, detail="Too many failed login attempts. Try again in 15 minutes.")
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        logger.warning(f"Rate limit exceeded for IP: {ip}")
        raise HTTPException(status_code=429, detail="Too many requests. Please try again later.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    credential = data.get("credential", "")
    if not credential:
        raise HTTPException(status_code=400, detail="Missing Google credential")
    try:
        idinfo = id_token.verify_oauth2_token(credential, google_requests.Request(), CLIENT_ID)
        email = idinfo.get("email")
        if not email:
            raise HTTPException(status_code=401, detail="Email not found in token")
        manifest = load_manifest()
        authorized = manifest.get("authorizedEmails", [])
        is_auth = email.lower() == "iancheung@berkeley.edu" or email.lower() in [e.lower() for e in authorized]
        if is_auth:
            logger.info(f"Successful login for {email} from IP: {ip}")
            login_fails[ip] = []
            token = create_session(ip, email.lower())
            response = JSONResponse(content={"success": True, "message": "Login successful", "email": email}, status_code=200)
            response.set_cookie(SESSION_COOKIE, token, max_age=SESSION_MAX_AGE, httponly=True, secure=True, samesite="strict")
            return response
        else:
            login_fails[ip].append(time.time())
            logger.warning(f"Unauthorized email {email} from IP: {ip}")
            raise HTTPException(status_code=401, detail="Email not authorized")
    except ValueError as e:
        login_fails[ip].append(time.time())
        logger.warning(f"Invalid Google token from IP: {ip}: {str(e)}")
        raise HTTPException(status_code=401, detail="Invalid token")

@app.post("/logout")
async def logout(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated")
    token = request.cookies.get(SESSION_COOKIE)
    if token:
        sessions.pop(token, None)
        logger.info(f"Session logged out: {token}")
    response = JSONResponse(content={"success": True, "message": "Logged out successfully"}, status_code=200)
    response.delete_cookie(SESSION_COOKIE)
    return response

@app.get("/", response_class=HTMLResponse)
async def root(request: Request):
    template_path = Path(__file__).parent / "templates" / "index.html"
    if template_path.exists():
        with open(template_path, "r") as f:
            return f.read()
    return "<h1>Template not found</h1>"

@app.post("/search")
async def search(request: Request, q: str):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        logger.warning(f"Search rate limit exceeded for IP: {ip}")
        raise HTTPException(status_code=429, detail="Too many search requests. Please slow down.")
    if not q or len(q.strip()) == 0:
        return {"results": {}, "query": q, "metadata": {}, "total_database_rows": sum(m.get("total_rows", 0) for m in meta.values())}
    results = search_dfs(q.strip())
    res = {}
    for name, rows in results.items():
        res[name] = {"rows": rows, "last_modified_iso": meta.get(name, {}).get("last_modified_iso", "Unknown"), "total_rows": meta.get(name, {}).get("total_rows", 0)}
    total = sum(m.get("total_rows", 0) for m in meta.values())
    return {"results": res, "query": q, "total_matches": sum(len(v["rows"]) for v in res.values()), "total_database_rows": total}

@app.get("/student-entries")
async def get_entries(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    entries = fetch_entries()
    return {"entries": entries, "total_count": len(entries), "success": len(entries) > 0}

@app.post("/update-entry-color")
async def set_color(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    row_idx = data.get("row_index")
    color = data.get("color_hex")
    if row_idx is None or not isinstance(row_idx, int) or row_idx < 1:
        raise HTTPException(status_code=400, detail="Invalid row index")
    if not color or not isinstance(color, str) or not color.startswith("#"):
        raise HTTPException(status_code=400, detail="Invalid color format")
    try:
        if not sheets:
            raise Exception("Google Sheets API not initialized")
        hex_c = color.lstrip("#")
        rgb = tuple(int(hex_c[i:i+2], 16) / 255 for i in (0, 2, 4))
        c = {"red": rgb[0], "green": rgb[1], "blue": rgb[2]}
        reqs = []
        for col_idx in range(1, 6):
            reqs.append({
                "updateCells": {
                    "range": {"sheetId": sid, "startRowIndex": row_idx, "endRowIndex": row_idx + 1, "startColumnIndex": col_idx, "endColumnIndex": col_idx + 1},
                    "rows": [{"values": [{"userEnteredFormat": {"backgroundColor": c}}]}],
                    "fields": "userEnteredFormat.backgroundColor"
                }
            })
        sheets.spreadsheets().batchUpdate(spreadsheetId=SHEET_ID, body={"requests": reqs}).execute()
        logger.info(f"Updated color for row {row_idx} to {color}")
        return {"success": True, "message": "Color updated successfully"}
    except Exception as e:
        logger.error(f"Error updating entry color: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to update color")

@app.post("/update-entry-notes")
async def set_notes(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    row_idx = data.get("row_index")
    text = data.get("notes_text", "")
    if row_idx is None or not isinstance(row_idx, int) or row_idx < 1:
        raise HTTPException(status_code=400, detail="Invalid row index")
    try:
        if not sheets:
            raise Exception("Google Sheets API not initialized")
        cell = f"'Form Responses 1'!H{row_idx + 1}"
        sheets.spreadsheets().values().update(
            spreadsheetId=SHEET_ID, range=cell, valueInputOption="USER_ENTERED",
            body={"values": [[text]]}
        ).execute()
        logger.info(f"Updated notes for row {row_idx}")
        return {"success": True, "message": "Notes updated successfully"}
    except Exception as e:
        logger.error(f"Error updating entry notes: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to update notes")

@app.post("/update-entry-discord-username")
async def set_discord_user(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    row_idx = data.get("row_index")
    username = data.get("discord_username", "")
    if row_idx is None or not isinstance(row_idx, int) or row_idx < 1:
        raise HTTPException(status_code=400, detail="Invalid row index")
    try:
        if not sheets:
            raise Exception("Google Sheets API not initialized")
        cell = f"'Form Responses 1'!F{row_idx + 1}"
        sheets.spreadsheets().values().update(
            spreadsheetId=SHEET_ID, range=cell, valueInputOption="USER_ENTERED",
            body={"values": [[username]]}
        ).execute()
        logger.info(f"Updated discord username for row {row_idx} to '{username}'")
        return {"success": True, "message": "Discord username updated successfully"}
    except Exception as e:
        logger.error(f"Error updating discord username: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to update discord username")

@app.post("/update-entry-discord-id")
async def set_discord_id(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    row_idx = data.get("row_index")
    did = data.get("discord_user_id", "")
    if row_idx is None or not isinstance(row_idx, int) or row_idx < 1:
        raise HTTPException(status_code=400, detail="Invalid row index")
    try:
        if not sheets:
            raise Exception("Google Sheets API not initialized")
        cell = f"'Form Responses 1'!I{row_idx + 1}"
        sheets.spreadsheets().values().update(
            spreadsheetId=SHEET_ID, range=cell, valueInputOption="USER_ENTERED",
            body={"values": [[did]]}
        ).execute()
        logger.info(f"Updated discord user ID for row {row_idx} to '{did}'")
        return {"success": True, "message": "Discord user ID updated successfully"}
    except Exception as e:
        logger.error(f"Error updating discord user ID: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to update discord user ID")


@app.get("/me")
async def current_user(request: Request):
    token = request.cookies.get(SESSION_COOKIE)
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    data = sessions.get(token)
    if not data:
        raise HTTPException(status_code=401, detail="Not authenticated")
    if time.time() > data.get("expires_at", 0):
        sessions.pop(token, None)
        raise HTTPException(status_code=401, detail="Session expired")
    return {"email": data.get("email", "")}

@app.get("/roles")
async def get_roles(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    manifest = load_manifest()
    return {"success": True, "emails": manifest.get("authorizedEmails", [])}

@app.post("/roles")
async def add_role(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    email = data.get("email", "").strip().lower()
    if not email or "@" not in email:
        raise HTTPException(status_code=400, detail="Invalid email address")
    if email == "iancheung@berkeley.edu":
        return {"success": True, "message": "Email is already permanently authorized"}
    manifest = load_manifest()
    emails = manifest.get("authorizedEmails", [])
    if email not in emails:
        emails.append(email)
        manifest["authorizedEmails"] = emails
        save_manifest(manifest)
        logger.info(f"Added authorized email: {email}")
    return {"success": True, "message": "Email authorized successfully"}

@app.delete("/roles")
async def remove_role(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    email = data.get("email", "").strip().lower()
    if not email:
        raise HTTPException(status_code=400, detail="Invalid email address")
    if email == "iancheung@berkeley.edu":
        raise HTTPException(status_code=403, detail="Cannot remove the master administrator")
    token = request.cookies.get(SESSION_COOKIE)
    session_data = sessions.get(token) if token else None
    if session_data:
        session_email = session_data.get("email", "").lower()
        if session_email == email:
            raise HTTPException(status_code=403, detail="You cannot remove yourself")
    manifest = load_manifest()
    emails = manifest.get("authorizedEmails", [])
    if email in emails:
        emails.remove(email)
        manifest["authorizedEmails"] = emails
        save_manifest(manifest)
        logger.info(f"Removed authorized email: {email}")
    async with sessions_lock:
        tokens_to_remove = [t for t, s in list(sessions.items()) if s.get("email") == email]
        for t in tokens_to_remove:
            sessions.pop(t, None)
    return {"success": True, "message": "Email removed successfully"}


@app.post("/discord/user-profile")
async def discord_profile(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    username = data.get("discord_username", "").strip().lower()
    if not username:
        raise HTTPException(status_code=400, detail="Discord username required")
    try:
        from bot import client
        if not client or not client.is_ready():
            logger.warning("Discord user-profile request but bot not ready")
            return {"success": False, "error": "Bot not ready", "message": "Discord bot is still connecting. Please try again in a moment."}
        member = await get_user(username)
        if not member:
            logger.info(f"Discord user not found: {username}")
            return {"success": False, "error": "User not found", "message": f"Could not find Discord user '{username}' in the server"}
        profile = await get_profile(member.id)
        if not profile:
            return {"success": False, "error": "Profile fetch failed", "message": "Failed to fetch user profile"}
        return {"success": True, "profile": profile}
    except Exception as e:
        logger.error(f"Error fetching Discord profile: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to fetch Discord profile")

@app.post("/discord/assign-role")
async def assign_discord_role(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    user_id = data.get("user_id")
    role_name = data.get("role_name", "").strip()
    send_welcome = data.get("send_welcome_msg", True)
    if not user_id:
        raise HTTPException(status_code=400, detail="Invalid user ID")
    try:
        user_id = int(user_id) if isinstance(user_id, str) else user_id
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="Invalid user ID format")
    valid_roles = [
        "Class of 2030", "Class of 2029", "Class of 2028", "Class of 2027",
        "Class of 2026", "Class of 2025", "Class of 2024", "Class of 2023",
        "Transfer Class of 2028", "Transfer Class of 2027", "Transfer Class of 2026",
        "Transfer Class of 2025", "Transfer Class of 2024",
        "Grad & Professional Student", "UCEAP"
    ]
    if role_name not in valid_roles:
        raise HTTPException(status_code=400, detail="Invalid role name")
    try:
        from bot import client
        if not client or not client.is_ready():
            logger.warning("Discord assign-role request but bot not ready")
            return {"success": False, "error": "Bot not ready", "message": "Discord bot is not connected. Role assignment unavailable."}
        success = await assign_role(user_id, role_name, send_welcome)
        if success:
            logger.info(f"Successfully assigned role '{role_name}' to user {user_id}")
            return {"success": True, "message": f"Role '{role_name}' assigned successfully"}
        else:
            logger.warning(f"Failed to assign role to user {user_id}")
            return {"success": False, "error": "Assignment failed", "message": "Failed to assign role to user"}
    except Exception as e:
        logger.error(f"Error assigning role: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to assign role")


@app.post("/link-sheet")
async def link_sheet(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    url = data.get("sheet_url", "").strip()
    if not url:
        raise HTTPException(status_code=400, detail="Sheet URL required")
    if "docs.google.com/spreadsheets" not in url:
        raise HTTPException(status_code=400, detail="Invalid Google Sheets URL")
    match = re.search(r'/spreadsheets/d/([a-zA-Z0-9-_]+)', url)
    if not match:
        raise HTTPException(status_code=400, detail="Invalid Google Sheets URL format. Expected: https://docs.google.com/spreadsheets/d/[SHEET_ID]/...")
    sheet_id = match.group(1)
    gid_match = re.search(r'[?#&]gid=(\d+)', url)
    gid = gid_match.group(1) if gid_match else "0"
    try:
        creds_file = Path(__file__).parent / "service-account-key.json"
        if not creds_file.exists():
            raise HTTPException(status_code=500, detail="Service account credentials not found")
        creds = Credentials.from_service_account_file(
            creds_file, scopes=['https://www.googleapis.com/auth/drive', 'https://www.googleapis.com/auth/spreadsheets']
        )
        drive = build('drive', 'v3', credentials=creds)
        ss = build('sheets', 'v4', credentials=creds)
        csv_content, sp_name, sname = download_sheet(drive, ss, sheet_id, gid)
        if not csv_content or not sname:
            raise HTTPException(status_code=400, detail="Failed to access sheet. Ensure it is shared with the service account (Viewer access) and is a native Google Sheet or XLSX file.")
        filename = f"{sp_name}.csv"
        data_dir = Path(__file__).parent / "data"
        data_dir.mkdir(exist_ok=True)
        out = data_dir / filename
        manifest = load_manifest()
        if out.exists():
            existing = manifest.get("sheetFiles", {}).get(filename)
            if existing and existing.get("sheetId") != sheet_id:
                counter = 2
                while True:
                    new_name = f"{sp_name} {counter}.csv"
                    new_path = data_dir / new_name
                    new_info = manifest.get("sheetFiles", {}).get(new_name)
                    if not new_path.exists() or (new_info and new_info.get("sheetId") == sheet_id):
                        filename = new_name
                        out = new_path
                        break
                    counter += 1
        with open(out, 'w') as f:
            f.write(csv_content)
        logger.info(f"Successfully linked sheet: {filename}")
        manifest = load_manifest()
        manifest["sheetFiles"][filename] = {
            "sheetId": sheet_id, "gid": gid, "sheetName": sname,
            "spreadsheetName": sp_name, "uploadedAt": datetime.now(tz=timezone.utc).isoformat()
        }
        save_manifest(manifest)
        reload_csvs()
        return {"success": True, "message": "Sheet linked successfully", "filename": filename, "sheetName": sname}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error linking sheet: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to link sheet")

@app.delete("/delete-sheet")
async def delete_sheet(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    filename = data.get("filename", "").strip()
    if not filename:
        raise HTTPException(status_code=400, detail="Filename required")
    if "/" in filename or "\\" in filename or filename.startswith("."):
        raise HTTPException(status_code=400, detail="Invalid filename")
    try:
        data_dir = Path(__file__).parent / "data"
        path = data_dir / filename
        if not path.exists():
            raise HTTPException(status_code=404, detail="File not found")
        path.unlink()
        logger.info(f"Unlinked sheet: {filename}")
        manifest = load_manifest()
        if filename in manifest.get("sheetFiles", {}):
            del manifest["sheetFiles"][filename]
            save_manifest(manifest)
        reload_csvs()
        return {"success": True, "message": "Sheet unlinked successfully", "filename": filename}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting CSV: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to unlink sheet")

@app.post("/rename-sheet")
async def rename_sheet(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    ip = client_ip(request)
    if rate_limited(ip, searches, SEARCH_MAX, SEARCH_WIN):
        raise HTTPException(status_code=429, detail="Too many requests. Please slow down.")
    try:
        data = await request.json()
    except:
        raise HTTPException(status_code=400, detail="Invalid request")
    filename = data.get("filename", "").strip()
    new_name = data.get("newName", "").strip()
    if not filename:
        raise HTTPException(status_code=400, detail="Filename required")
    if not new_name:
        raise HTTPException(status_code=400, detail="New name required")
    if "/" in filename or "\\" in filename or filename.startswith("."):
        raise HTTPException(status_code=400, detail="Invalid filename")
    if "/" in new_name or "\\" in new_name or new_name.startswith("."):
        raise HTTPException(status_code=400, detail="Invalid new name")
    try:
        data_dir = Path(__file__).parent / "data"
        path = data_dir / filename
        new_filename = f"{new_name}.csv"
        new_path = data_dir / new_filename
        if not path.exists():
            raise HTTPException(status_code=404, detail="File not found")
        if new_path.exists() and new_filename != filename:
            raise HTTPException(status_code=400, detail="A file with that name already exists")
        path.rename(new_path)
        logger.info(f"Renamed sheet: {filename} -> {new_filename}")
        manifest = load_manifest()
        if filename in manifest.get("sheetFiles", {}):
            fm = manifest["sheetFiles"].pop(filename)
            manifest["sheetFiles"][new_filename] = fm
            save_manifest(manifest)
        reload_csvs()
        return {"success": True, "message": "Sheet renamed successfully", "oldFilename": filename, "newFilename": new_filename}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error renaming CSV: {str(e)}")
        raise HTTPException(status_code=500, detail="Failed to rename sheet")

@app.get("/sheet-manifest")
async def sheet_manifest(request: Request):
    if not await authed(request):
        raise HTTPException(status_code=401, detail="Not authenticated. Please login first.")
    manifest = load_manifest()
    return {"success": True, "sheetFiles": manifest.get("sheetFiles", {}), "total": len(manifest.get("sheetFiles", {}))}

uvicorn.run(app, host="0.0.0.0", port=4187)
